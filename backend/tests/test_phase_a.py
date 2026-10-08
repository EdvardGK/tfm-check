"""Phase A (HI90): codes composed from parts, the shape check per segment,
roles found in the model, status from MMI, and the register.
Run from backend/: python -m pytest tests -q"""

from __future__ import annotations

import io
import sys
from pathlib import Path

import ifcopenshell
import ifcopenshell.api
import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from engine.ifc_io import detect_discipline_from_filename, list_products  # noqa: E402
from engine.inventory import ModelIndex, inventory_payload, phase_of, preview, values_payload  # noqa: E402
from engine.presets import PRESETS, preset_to_rules_dict  # noqa: E402
from engine.register import build_register_xlsx, register_rows  # noqa: E402
from engine.rules import TFMRules, compose  # noqa: E402
from engine.shape import diagnose  # noqa: E402
from engine.checks import run_checks  # noqa: E402

SET = "P_Info"
LOK, SYS, KOMP, MMI = "TFM PLASSERINGS-ID", "TFM SYSTEMFOREKOMST-ID", "TFM KOMPONENTTYPE-ID", "P_MMI"


def statsbygg(**kw) -> TFMRules:
    d = preset_to_rules_dict(PRESETS[0])
    d.update(kw)
    return TFMRules.from_dict(d)


# ---- shape -----------------------------------------------------------------

@pytest.mark.parametrize("code,ok,fix", [
    ("+02=320.003-JP401", True, ""),
    ("+02=360.017", True, ""),                    # system without a component
    ("+02=320.3-JP401", False, "+02=320.003-JP401"),   # running number padded
    ("+02=320-003-JP401", False, "+02=320.003-JP401"),  # separator replaced
    (" +02=320.003-jp401", False, "+02=320.003-JP401"),  # trimmed, upper-cased
    ("+05.=360.001-JP4", False, "+05=360.001-JP004"),    # trailing separator
    ("+04=360.001-SQ.001T", False, ""),          # T is content: no fix
    ("+04=360.001-KRA.016", False, ""),          # three letters: no fix
    ("+4=360.001", False, ""),                    # Lokasjon never padded
    ("+02=32.003", False, ""),                    # Systemkode never padded
    ("-SQ.001T", False, ""),
])
def test_diagnose(code, ok, fix):
    d = diagnose(code, statsbygg(part_digits={"lokasjon": 2}))
    assert d.ok is ok
    assert d.fix == fix
    if not ok:
        assert d.reason


def test_diagnose_reasons_name_the_segment():
    r = statsbygg(part_digits={"lokasjon": 2})
    assert diagnose("+04=360.001-KRA.016", r).reason.startswith("Komponent «KRA», skal være 2 bokstaver")
    assert diagnose("+04-SQ.001T", r).reason == "mangler «=» Systemkode"
    assert "6 tegn" in diagnose("+04=360.001", statsbygg()).reason


def test_compose():
    assert compose({"lokasjon": "04", "system": "360.001", "komponent": "JP401"}) == "+04=360.001-JP401"
    assert compose({"lokasjon": "+04", "system": "", "komponent": "SQ.001T"}) == "+04-SQ.001T"
    assert compose({}) == ""


@pytest.mark.parametrize("name,fag", [
    ("HI90_RIV.ifc", "RIV"), ("HI90_RIV_MMI700.ifc", "RIV"), ("ST28_RIE.ifc", "RIE"),
    ("KNM_RIA.ifc", "RIA"), ("HI90_ARK-A.ifc", "ARK"), ("modell.ifc", None),
])
def test_fag_from_file_name(name, fag):
    assert detect_discipline_from_filename(name) == fag


@pytest.mark.parametrize("mmi,phase", [("300", "ny"), ("700", "bevares"), ("875", "ombruk"), ("980", "rives"), ("x", "")])
def test_phase_of(mmi, phase):
    assert phase_of(mmi) == phase


# ---- a model with the code in parts ----------------------------------------

def _model():
    f = ifcopenshell.file(schema="IFC4")
    project = ifcopenshell.api.run("root.create_entity", f, ifc_class="IfcProject", name="P")
    ifcopenshell.api.run("unit.assign_unit", f)
    site = ifcopenshell.api.run("root.create_entity", f, ifc_class="IfcSite", name="S")
    building = ifcopenshell.api.run("root.create_entity", f, ifc_class="IfcBuilding", name="B")
    ifcopenshell.api.run("aggregate.assign_object", f, products=[site], relating_object=project)
    ifcopenshell.api.run("aggregate.assign_object", f, products=[building], relating_object=site)
    st = ifcopenshell.api.run("root.create_entity", f, ifc_class="IfcBuildingStorey", name="4. etasje")
    ifcopenshell.api.run("aggregate.assign_object", f, products=[st], relating_object=building)
    rows = [
        ("02", "320.003", "JP401", "300"),
        ("02", "360.017", "", "875"),
        ("04", "360.001", "SQ.001T", "700"),
        ("05.", "360.001", "JP4", "950"),
        ("", "", "", "300"),
    ]
    for i, (lok, sys_, komp, mmi) in enumerate(rows):
        el = ifcopenshell.api.run("root.create_entity", f, ifc_class="IfcPump", name=f"E{i}")
        ifcopenshell.api.run("spatial.assign_container", f, products=[el], relating_structure=st)
        pset = ifcopenshell.api.run("pset.add_pset", f, product=el, name=SET)
        ifcopenshell.api.run("pset.edit_pset", f, pset=pset, properties={
            LOK: lok, SYS: sys_, KOMP: komp, MMI: mmi, "TFM-ID": "",
        })
    return f


@pytest.fixture(scope="module")
def model():
    f = _model()
    products = list_products(f)
    return f, products, ModelIndex(f, products)


def parts_rules(**kw) -> TFMRules:
    return statsbygg(
        part_digits={"lokasjon": 2},
        tfm_mode="parts",
        tfm_parts={"lokasjon": ["pset", SET, LOK], "system": ["pset", SET, SYS], "komponent": ["pset", SET, KOMP]},
        status_location=["pset", SET, MMI],
        **kw,
    )


def test_roles_found_in_the_model(model):
    _, _, index = model
    inv = inventory_payload(index, [])
    roles = inv["roles"]
    assert roles["lokasjon"]["standard"]["n"] == 0
    assert roles["lokasjon"]["candidates"][0]["location"] == ["pset", SET, LOK]
    assert roles["system"]["candidates"][0]["location"] == ["pset", SET, SYS]
    assert roles["komponent"]["candidates"][0]["location"] == ["pset", SET, KOMP]
    assert roles["status"]["candidates"][0]["location"] == ["pset", SET, MMI]


def test_composed_preview(model):
    _, _, index = model
    p = preview(index, parts_rules())
    assert p["valued"] == 4
    assert p["matched"] == 2
    off = {o["v"]: o for o in p["off"]}
    assert off["+05.=360.001-JP4"]["fix"] == "+05=360.001-JP004"
    assert off["+04=360.001-SQ.001T"]["fix"] == ""
    assert off["+04=360.001-SQ.001T"]["reason"]


def test_status_values(model):
    _, _, index = model
    v = values_payload(index, ["pset", SET, MMI])
    assert v["phases"] == {"ny": 2, "bevares": 1, "ombruk": 1, "rives": 1, "": 0}


def test_run_checks_on_composed_codes(model):
    f, products, index = model
    rules = parts_rules()
    res = run_checks(f, products, rules, {}, {}, index.code_values(rules))
    assert res["checks"]["has_code"]["n"] == 2
    assert res["checks"]["has_code"]["total"] == 5


def test_register(model):
    f, products, index = model
    columns, rows, summary, roll = register_rows(f, products, index, parts_rules(), "HI90_RIV.ifc", "RIV")
    assert columns[:4] == ["Kode som funnet", LOK, SYS, KOMP]
    assert len(rows) == 4
    by = {r["Kode som funnet"]: r for r in rows}
    r = by["+02=320.003-JP401"]
    assert (r["Gyldig format"], r["Status"], r["MMI"], r["Systemkode"], r["Fag"]) == ("ja", "ny", "300", "320", "RIV")
    assert r["Etasje (modell)"] == "4. etasje" and r["IFC-klasse"] == "IfcPump"
    assert by["+05.=360.001-JP4"]["Ny kode"] == "+05=360.001-JP004"
    assert by["+05.=360.001-JP4"][LOK] == "05."
    assert by["+05.=360.001-JP4"]["Status"] == "rives"
    assert summary["Gyldig format"] == 2 and summary["Med ny kode"] == 1
    data = build_register_xlsx(columns, rows, summary, roll)
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Sammendrag", "Register", "Systemkoder", "Komponentkoder"]
    ws = wb["Register"]
    assert [c.value for c in ws[1]][:4] == columns[:4]
    assert ws.max_row == 5
