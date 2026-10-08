"""The TFM register: one row per object carrying a code, as an .xlsx.

Each row: the code as found (verbatim, and per source field when the code is
composed), whether it takes the format and why not, whether its linked parts
are real entries of their standard, the mechanically fixed code, the object's
status from its MMI, and its link to the model (file, GlobalId, IFC class,
storey, systems). Two more sheets roll the system and component codes up.
"""

from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .codes import coded_objects, rollup
from .constants import sequence_to_example
from .inventory import ModelIndex, phase_of
from .rules import ASPECTS, TFMRules

KODE = "Kode som funnet"
COLUMNS_AFTER = [
    "Kilde", "Gyldig format", "Gyldig i standard", "Årsak", "Ny kode", "Status", "MMI",
    "Modellkobling (fil)", "GlobalId", "IFC-klasse", "Etasje (modell)",
    "Fag", "Systemkode", "Beskrivelse system", "Komponentkode", "Beskrivelse komponent",
    "IfcSystem", "Avvik",
]
SYS_COLUMNS = ["Systemkode", "Beskrivelse", "Gyldig i standard", "Årsak", "Objekter", "Systemer", "Komponenter"]
COMP_COLUMNS = ["Komponentkode", "Beskrivelse", "Gyldig i standard", "Årsak", "Objekter", "Systemkoder"]

# Phase colours (design canon 2026-10-08, HI90 Typebank): N, EK, R, EW.
PHASE_FILL = {"ny": "FAF050", "bevares": "D9A633", "ombruk": "66E6B3", "rives": "803366"}
PHASE_FONT = {"rives": "FFFFFF"}
OK_FILL, BAD_FILL = "C6EFCE", "FFC7CE"
HEAD_FILL = "D9D9D9"


def _yn(v: bool | None) -> str:
    return "" if v is None else ("ja" if v else "nei")


def _loc_text(loc) -> str:
    if not loc:
        return ""
    kind, pset, prop = (list(loc) + [None, None, None])[:3]
    if kind == "pset":
        return f"{pset}.{prop}"
    return prop or ""


def _field_columns(rules: TFMRules) -> list[tuple[str, tuple]]:
    """(header, source) per source field of a composed code, aspect order,
    each source once. The header is the property's own name."""
    if not rules.composed:
        return []
    out, seen = [], set()
    for a in ASPECTS:
        loc = rules.tfm_parts.get(a)
        if not loc or loc in seen:
            continue
        seen.add(loc)
        out.append((loc[2] or _loc_text(loc), loc))
    names: dict[str, int] = {}
    for h, _ in out:
        names[h] = names.get(h, 0) + 1
    return [(_loc_text(loc) if names[h] > 1 else h, loc) for h, loc in out]


def _storey_names(ifc) -> dict[int, str]:
    out = {}
    for st in ifc.by_type("IfcBuildingStorey"):
        for rel in getattr(st, "ContainsElements", None) or []:
            for el in rel.RelatedElements:
                out[el.id()] = st.Name or ""
    return out


def _systems(ifc) -> dict[int, list[str]]:
    out: dict[int, list[str]] = {}
    try:
        groups = ifc.by_type("IfcRelAssignsToGroup")
    except RuntimeError:
        return out
    for rel in groups:
        g = rel.RelatingGroup
        if g is None or not g.is_a("IfcSystem"):
            continue
        for o in rel.RelatedObjects or []:
            out.setdefault(o.id(), []).append(g.Name or "")
    return out


def _counts(items: list[dict]) -> str:
    return ", ".join(f"{x['v']} ({x['n']})" for x in items)


def register_rows(ifc, products, index: ModelIndex, rules: TFMRules,
                  file_name: str, fag: str | None) -> tuple[list[str], list[dict], dict, dict]:
    """(columns, rows, summary, rollup) for one model."""
    fields = _field_columns(rules)
    columns = [KODE, *[h for h, _ in fields], *COLUMNS_AFTER]
    field_values = [(h, index.values_for(loc) or {}) for h, loc in fields]
    status = (index.values_for(rules.status_location) if rules.status_location else None) or {}
    storeys = _storey_names(ifc)
    systems = _systems(ifc)
    by_id = {e.id(): e for e in products}

    coded, n_excluded = coded_objects(index, rules)
    rows: list[dict] = []
    for c in coded:
        e = by_id.get(c.pid)
        if e is None:
            continue
        mmi = status.get(c.pid, "")
        sv, kv = c.validity.get("systemkode"), c.validity.get("komponent")
        row = {KODE: c.code}
        for h, vals in field_values:
            row[h] = vals.get(c.pid, "")
        row.update({
            "Kilde": "modell",
            "Gyldig format": "ja" if c.diag.ok else "nei",
            "Gyldig i standard": _yn(c.std_ok),
            "Årsak": c.reason(),
            "Ny kode": c.diag.fix,
            "Status": phase_of(mmi),
            "MMI": mmi,
            "Modellkobling (fil)": file_name,
            "GlobalId": e.GlobalId,
            "IFC-klasse": e.is_a(),
            "Etasje (modell)": storeys.get(c.pid, ""),
            "Fag": fag or "",
            "Systemkode": c.parts.get("systemkode", ""),
            "Beskrivelse system": sv.description if sv else "",
            "Komponentkode": c.parts.get("komponent", ""),
            "Beskrivelse komponent": kv.description if kv else "",
            "IfcSystem": "; ".join(sorted(set(systems.get(c.pid, [])))),
            "Avvik": "",
        })
        rows.append(row)

    rows.sort(key=lambda r: (r[KODE], r["GlobalId"]))
    roll = rollup(coded, rules)
    summary = {
        "Fil": file_name,
        "Fag": fag or "",
        "Kilde": (" + ".join(_loc_text(rules.tfm_parts.get(a)) for a in ASPECTS if rules.tfm_parts.get(a))
                  if rules.composed else _loc_text(rules.tfm_location)),
        "Format": "  |  ".join(sequence_to_example(p.get("sequence") or []) for p in rules.patterns),
        "Standard": ", ".join(f"{p}: {lbl}" for p, lbl in roll["links"].items()),
        "Status fra": _loc_text(rules.status_location),
        "Elementer": len(index.product_ids),
        "Med kode": len(rows),
        "Utenfor scope": n_excluded,
        "Gyldig format": sum(1 for r in rows if r["Gyldig format"] == "ja"),
        "Ugyldig format": sum(1 for r in rows if r["Gyldig format"] == "nei"),
        "Gyldig i standard": sum(1 for r in rows if r["Gyldig i standard"] == "ja"),
        "Ugyldig i standard": sum(1 for r in rows if r["Gyldig i standard"] == "nei"),
        "Med ny kode": sum(1 for r in rows if r["Ny kode"]),
        "Dato": date.today().isoformat(),
    }
    for ph in ("ny", "bevares", "ombruk", "rives"):
        summary[f"Status {ph}"] = sum(1 for r in rows if r["Status"] == ph)
    return columns, rows, summary, roll


def _head(ws, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=HEAD_FILL)
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"


def _fill_yn(ws, row: int, col: int, value: str) -> None:
    if value in ("ja", "nei"):
        ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=OK_FILL if value == "ja" else BAD_FILL)


def _table(wb, title: str, columns: list[str], rows: list[dict]) -> None:
    ws = wb.create_sheet(title)
    ws.append(columns)
    yn_cols = [i + 1 for i, c in enumerate(columns) if c in ("Gyldig format", "Gyldig i standard")]
    for r in rows:
        ws.append([r.get(c, "") for c in columns])
        for col in yn_cols:
            _fill_yn(ws, ws.max_row, col, r.get(columns[col - 1], ""))
    _head(ws, len(columns))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(1, ws.max_row)}"
    for idx, c in enumerate(columns, start=1):
        longest = max([len(str(c))] + [len(str(r.get(c, ""))) for r in rows[:500]])
        ws.column_dimensions[get_column_letter(idx)].width = min(60, max(10, longest + 2))
    return ws


def build_register_xlsx(columns: list[str], rows: list[dict], summary: dict, roll: dict | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sammendrag"
    ws.append(["Felt", "Verdi"])
    for k, v in summary.items():
        ws.append([k, v])
    _head(ws, 2)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 60

    reg = _table(wb, "Register", columns, rows)
    st_col = columns.index("Status") + 1
    for i, r in enumerate(rows, start=2):
        ph = r.get("Status")
        if ph in PHASE_FILL:
            st = reg.cell(row=i, column=st_col)
            st.fill = PatternFill("solid", fgColor=PHASE_FILL[ph])
            if ph in PHASE_FONT:
                st.font = Font(color=PHASE_FONT[ph])

    if roll is not None:
        _table(wb, "Systemkoder", SYS_COLUMNS, [{
            "Systemkode": s["code"], "Beskrivelse": s["description"], "Gyldig i standard": _yn(s["valid"]),
            "Årsak": s["reason"], "Objekter": s["n"], "Systemer": _counts(s["systems"]),
            "Komponenter": _counts(s["components"]),
        } for s in roll["systems"]])
        _table(wb, "Komponentkoder", COMP_COLUMNS, [{
            "Komponentkode": c["code"], "Beskrivelse": c["description"], "Gyldig i standard": _yn(c["valid"]),
            "Årsak": c["reason"], "Objekter": c["n"], "Systemkoder": _counts(c["systems"]),
        } for c in roll["components"]])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
